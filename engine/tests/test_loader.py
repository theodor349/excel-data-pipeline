import json
from unittest.mock import MagicMock, patch

import openpyxl
import polars as pl
import pytest
from openpyxl.worksheet.table import Table

from engine.loader import _read_csv_path, _read_jsonl_path, read_csv, read_excel, read_jsonl, read_sql


def _write_config(path, data):
    path.write_text(json.dumps(data), encoding="utf-8")


# ---------------------------------------------------------------------------
# _read_jsonl_path — no config involved
# ---------------------------------------------------------------------------


def test_read_jsonl_path_reads_records(tmp_path):
    jsonl_file = tmp_path / "data.jsonl"
    jsonl_file.write_text('{"a": 1, "b": "x"}\n{"a": 2, "b": "y"}\n', encoding="utf-8")
    df = _read_jsonl_path(jsonl_file)
    assert list(df.columns) == ["a", "b"]
    assert len(df) == 2
    assert df["a"].to_list() == [1, 2]


def test_read_jsonl_path_skips_blank_lines(tmp_path):
    jsonl_file = tmp_path / "data.jsonl"
    jsonl_file.write_text('{"a": 1}\n\n{"a": 2}\n', encoding="utf-8")
    df = _read_jsonl_path(jsonl_file)
    assert len(df) == 2


def test_read_jsonl_path_empty_file(tmp_path):
    jsonl_file = tmp_path / "empty.jsonl"
    jsonl_file.write_text("", encoding="utf-8")
    df = _read_jsonl_path(jsonl_file)
    assert len(df) == 0


# ---------------------------------------------------------------------------
# read_jsonl
# ---------------------------------------------------------------------------


def test_read_jsonl_returns_expected_rows_and_columns(tmp_path, monkeypatch):
    jsonl_file = tmp_path / "events.jsonl"
    jsonl_file.write_text('{"id": 1, "type": "click"}\n{"id": 2, "type": "view"}\n', encoding="utf-8")

    config = {"jsonl_sources": {"events": str(jsonl_file)}, "csv_sources": {}, "excel_sources": {}, "connections": {}}
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    df = read_jsonl("events")
    assert list(df.columns) == ["id", "type"]
    assert len(df) == 2
    assert df["type"].to_list() == ["click", "view"]


def test_read_jsonl_raises_key_error_for_unknown_source(tmp_path, monkeypatch):
    config = {"jsonl_sources": {}, "csv_sources": {}, "excel_sources": {}, "connections": {}}
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    with pytest.raises(KeyError) as exc_info:
        read_jsonl("no_such_source")

    error_text = str(exc_info.value)
    assert "jsonl_sources" in error_text
    assert "no_such_source" in error_text


# ---------------------------------------------------------------------------
# _read_csv_path — no config involved
# ---------------------------------------------------------------------------


def test_read_csv_path_reads_file(tmp_path):
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("x,y\n1,2\n3,4\n")
    df = _read_csv_path(csv_file)
    assert list(df.columns) == ["x", "y"]
    assert len(df) == 2
    assert df["x"].to_list() == [1, 3]


# ---------------------------------------------------------------------------
# read_csv
# ---------------------------------------------------------------------------


def test_read_csv_returns_expected_rows_and_columns(tmp_path, monkeypatch):
    csv_file = tmp_path / "rates.csv"
    csv_file.write_text("currency,rate\nUSD,1.0\nEUR,0.9\n")

    config = {"csv_sources": {"fx_rates": str(csv_file)}, "excel_sources": {}, "connections": {}}
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    df = read_csv("fx_rates")
    assert list(df.columns) == ["currency", "rate"]
    assert len(df) == 2
    assert df["currency"].to_list() == ["USD", "EUR"]


def test_read_csv_raises_key_error_for_unknown_source(tmp_path, monkeypatch):
    config = {"csv_sources": {}, "excel_sources": {}, "connections": {}}
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    with pytest.raises(KeyError) as exc_info:
        read_csv("no_such_source")

    error_text = str(exc_info.value)
    assert "csv_sources" in error_text
    assert "no_such_source" in error_text


# ---------------------------------------------------------------------------
# read_excel — sheet by name
# ---------------------------------------------------------------------------


def _make_simple_xlsx(path, sheet_name="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["name", "value"])
    ws.append(["alpha", 10])
    ws.append(["beta", 20])
    wb.save(path)


def test_read_excel_reads_sheet_by_name(tmp_path, monkeypatch):
    xlsx_file = tmp_path / "data.xlsx"
    _make_simple_xlsx(xlsx_file, sheet_name="MySheet")

    config = {
        "excel_sources": {"my_source": str(xlsx_file)},
        "csv_sources": {},
        "connections": {},
    }
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    df = read_excel("my_source", "MySheet")
    assert list(df.columns) == ["name", "value"]
    assert len(df) == 2
    assert df["name"].to_list() == ["alpha", "beta"]


# ---------------------------------------------------------------------------
# read_excel — table by table name
# ---------------------------------------------------------------------------


def _make_xlsx_with_table(path, table_name="SalesTable"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["product", "amount"])
    ws.append(["widget", 100])
    ws.append(["gadget", 200])
    table = Table(displayName=table_name, ref="A1:B3")
    ws.add_table(table)
    wb.save(path)


def test_read_excel_reads_table_by_name_when_no_sheet_matches(tmp_path, monkeypatch):
    xlsx_file = tmp_path / "sales.xlsx"
    _make_xlsx_with_table(xlsx_file, table_name="SalesTable")

    config = {
        "excel_sources": {"sales": str(xlsx_file)},
        "csv_sources": {},
        "connections": {},
    }
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    df = read_excel("sales", "SalesTable")
    assert list(df.columns) == ["product", "amount"]
    assert len(df) == 2
    assert df["product"].to_list() == ["widget", "gadget"]


# ---------------------------------------------------------------------------
# read_excel — neither sheet nor table found
# ---------------------------------------------------------------------------


def test_read_excel_raises_when_neither_sheet_nor_table_found(tmp_path, monkeypatch):
    xlsx_file = tmp_path / "data.xlsx"
    _make_simple_xlsx(xlsx_file, sheet_name="RealSheet")

    config = {
        "excel_sources": {"my_source": str(xlsx_file)},
        "csv_sources": {},
        "connections": {},
    }
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    with pytest.raises(KeyError) as exc_info:
        read_excel("my_source", "NonExistent")

    assert "NonExistent" in str(exc_info.value)


# ---------------------------------------------------------------------------
# read_sql — mocked pyodbc
# ---------------------------------------------------------------------------


def test_read_sql_builds_connection_string_and_calls_read_sql(tmp_path, monkeypatch):
    config = {
        "connections": {
            "sales_db": {
                "type": "mssql",
                "host": "my-server",
                "database": "SalesDB",
                "username": "reader",
                "password": "secret",
            }
        },
        "excel_sources": {},
        "csv_sources": {},
    }
    config_path = tmp_path / "config.json"
    _write_config(config_path, config)
    monkeypatch.setenv("PIPELINE_CONFIG", str(config_path))

    mock_cursor = MagicMock()
    mock_cursor.description = [("id",), ("val",)]
    mock_cursor.fetchall.return_value = [(1, "a"), (2, "b")]

    mock_conn = MagicMock()
    mock_conn.cursor.return_value = mock_cursor

    mock_pyodbc = MagicMock()
    mock_pyodbc.connect.return_value = mock_conn

    with patch.dict("sys.modules", {"pyodbc": mock_pyodbc}):
        query = "SELECT id, val FROM some_table"
        result = read_sql("sales_db", query)

        connect_call_args = mock_pyodbc.connect.call_args[0][0]
        assert "my-server" in connect_call_args
        assert "SalesDB" in connect_call_args
        assert "reader" in connect_call_args

        mock_cursor.execute.assert_called_once_with(query)
        assert list(result.columns) == ["id", "val"]
        assert len(result) == 2
        assert result["id"].to_list() == [1, 2]
        mock_conn.close.assert_called_once()
